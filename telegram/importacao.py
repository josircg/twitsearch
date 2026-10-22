import csv
import re

from io import StringIO

from django.db import transaction
from openpyxl import load_workbook

from core import log_message
from .models import Canal, Lista

MAX_LINHAS = 5000
CABECALHOS = {'username', 'canal', 'channel', 'url', 'link', 'id'}
USERNAME_RE = re.compile(r'^[a-z][a-z0-9_]{3,31}$')


def normaliza_username(valor) -> str | None:
    """Converte URL, @username ou username em username minúsculo; None se inválido."""
    valor = str(valor or '').strip()
    valor = re.sub(r'^(https?://)?(www\.)?(t|telegram)\.me/', '', valor, flags=re.IGNORECASE)
    valor = valor.split('?')[0].split('/')[0]  # ignora links de mensagem e parâmetros
    username = valor.lstrip('@').lower()
    return username if USERNAME_RE.match(username) else None


def le_primeira_coluna(arquivo) -> list:
    """Lê a primeira coluna de um CSV ou XLSX, sem linhas vazias nem cabeçalho opcional."""
    if arquivo.name.lower().endswith('.csv'):
        texto = arquivo.read().decode('utf-8-sig')
        valores = [linha[0] for linha in csv.reader(StringIO(texto)) if linha]
    else:
        planilha = load_workbook(arquivo, read_only=True).active
        valores = [linha[0] for linha in planilha.iter_rows(values_only=True) if linha]
    valores = [str(valor).strip() for valor in valores if valor is not None and str(valor).strip()]
    if valores and valores[0].lower() in CABECALHOS:
        valores = valores[1:]
    return valores


def importa_lista(usuario, nome: str, publica: bool, valores: list) -> dict:
    """Cria a lista e seus canais de forma atômica. Retorna os totais da importação."""
    usernames = []
    ignorados = 0
    for valor in valores:
        username = normaliza_username(valor)
        if username:
            if username not in usernames:
                usernames.append(username)
        else:
            ignorados += 1
    if not usernames:
        raise ValueError('Nenhum canal válido foi encontrado no arquivo.')

    with transaction.atomic():
        lista = Lista.objects.create(nome=nome, dono=usuario, publica=publica)
        novos = 0
        for username in usernames:
            canal, criado = Canal.objects.get_or_create(username=username)
            if criado:
                novos += 1
                log_message(canal, f'Importado via formulário de projeto / {lista.nome}', user=usuario)
            lista.canais.add(canal)
    return {'lista': lista, 'lidos': len(valores), 'novos': novos, 'ignorados': ignorados}
